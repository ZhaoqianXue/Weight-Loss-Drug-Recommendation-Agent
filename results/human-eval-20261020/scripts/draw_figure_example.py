"""Draw the example answer for the §3 system figure under the rule frozen in figure_example_rule.md.

Pool: questions whose AskGLP1 answer (agent-0.8.5 first run) has consensus "No error" and "All supported", answer_type
"answered", at least one [cN] citation tag in its answer text, and verified = true (amendment 2026-10-23T03:55Z).
Draw: random.Random(20261023).choice over the eligible question ids sorted as strings.

    python results/human-eval-20261020/scripts/draw_figure_example.py --xlsx <export with a Consensus tab> [--out <file>]
"""
import argparse
import hashlib
import json
import random
import re
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[3]
STUDY = ROOT / 'results/human-eval-20261020'
KEY = ROOT / '.cache/human-eval-20261020/sheet/key.json'
RUN = ROOT / '.cache/human-eval-20261020/runs/final-askglp1-h120'
OUT = ROOT / '.cache/human-eval-20261020/figure-example/selection.json'
SEED = 20261023


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def consensus_labels(xlsx, key):
    """(question id, letter) -> (q1, q2) from the Consensus tab, laid out like a rater tab."""
    wb = load_workbook(xlsx, read_only=True)
    if 'Consensus' not in wb.sheetnames:
        raise SystemExit('no Consensus tab: the rule applies to consensus labels only')
    out, qid = {}, None
    for row in wb['Consensus'].iter_rows(values_only=True):
        if not row or len(row) < 9:
            continue
        if row[1]:
            qid = row[1]
        if row[2] in ('A', 'B', 'C') and qid in key['letters']:
            out[(qid, row[2])] = (row[7], row[8])
    return out


def cited_numbers(qid):
    text = (RUN / f'{qid}.md').read_text(encoding='utf-8')
    answer = text.split('# Answer', 1)[1].split('## Citations', 1)[0]
    return len(re.findall(r'\[c\d+\]', answer))


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--xlsx', required=True)
    ap.add_argument('--out', default=str(OUT))
    a = ap.parse_args(argv)
    key = json.loads(KEY.read_text())
    labels = consensus_labels(a.xlsx, key)
    runs = {r['id']: r for r in map(json.loads, (RUN / 'results.jsonl').read_text().splitlines())}
    pool = []
    for qid, letters in key['letters'].items():
        letter = next(l for l, s in letters.items() if s == 'askglp1')
        q1, q2 = labels.get((qid, letter), (None, None))
        if (q1 == 'No error' and q2 == 'All supported' and runs[qid]['answer_type'] == 'answered' and cited_numbers(qid) > 0
                and runs[qid]['verified'] is True):
            pool.append(qid)
    pool.sort()
    draw = random.Random(SEED).choice(pool) if pool else None
    result = {'seed': SEED, 'pool_size': len(pool), 'pool': pool, 'draw': draw,
              'rule_sha256': sha256(STUDY / 'figure_example_rule.md'), 'script_sha256': sha256(__file__),
              'consensus_workbook_sha256': sha256(a.xlsx)}
    Path(a.out).parent.mkdir(parents=True, exist_ok=True)
    Path(a.out).write_text(json.dumps(result, indent=1))
    print(json.dumps(result, indent=1))


if __name__ == '__main__':
    main()
