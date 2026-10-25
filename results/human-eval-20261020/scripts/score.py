"""Score the rating workbook of the RQ3 human evaluation.

Reads the workbook exported from Google Sheets (xlsx) and the local key (make_sheet.py), drops the yellow example rows,
and reports per rater tab, and for the consensus labels when a "Consensus" tab exists:
- per system, the share of answers meeting each criterion with a Wilson 95% interval: correct = no major error (and,
  stricter, no error); supported = all claims supported, among answers that make claims about the data;
- abstention when evidence is missing: the correct share on the 20 questions to decline or limit (raters judge those
  answers against a reference that says to decline or limit), together with an automatic check of each system's answer
  type against the question's expected behaviour (declined or limited on those 20; fully answered on the 100 others);
- AskGLP1 against B1 and against B2, paired by question: exact McNemar test and the paired difference in percentage
  points with a Newcombe (method 10) 95% interval;
- agreement within each rater pair: Cohen's kappa (linear-weighted for the three-level correctness scale).
The FAERS questions (area rq2) are also reported on their own, because the baselines have no FAERS data.

    python results/human-eval-20261020/scripts/score.py --xlsx <export.xlsx> [--out <dir>]
"""
import argparse
import json
import math
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[3]
KEY = ROOT / '.cache/human-eval-20261020/sheet/key.json'
QUESTIONS = ROOT / 'results/human-eval-20261020/eval_questions_human120.json'
SYSTEMS = ('askglp1', 'b1', 'b2')
Q1 = ['No error', 'Minor error', 'Major error']
Q2 = ['All supported', 'Some unsupported', 'None supported', 'No data claims']
RUNS = {'askglp1': 'final-askglp1-h120', 'b1': 'final-b1-h120', 'b2': 'final-b2-h120'}


def wilson(k, n):
    if not n:
        return None, None
    p, z = k / n, 1.96
    c = (p + z * z / (2 * n)) / (1 + z * z / n)
    h = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / (1 + z * z / n)
    return round(100 * max(0.0, c - h), 1), round(100 * min(1.0, c + h), 1)


def mcnemar_exact(b, c):
    """Two-sided exact McNemar p for discordant counts b (first only) and c (second only)."""
    n = b + c
    if n == 0:
        return 1.0
    k = min(b, c)
    p = sum(math.comb(n, i) for i in range(k + 1)) / 2 ** n
    return min(1.0, 2 * p)


def newcombe_paired(a, b, c, d):
    """Difference p1 - p2 (percentage points) for paired binary data with Newcombe's method 10 interval.
    a = both yes, b = first yes only, c = second yes only, d = both no."""
    n = a + b + c + d
    if not n:
        return None
    p1, p2 = (a + b) / n, (a + c) / n
    l1, u1 = [x / 100 for x in wilson(a + b, n)]
    l2, u2 = [x / 100 for x in wilson(a + c, n)]
    denom = math.sqrt((a + b) * (c + d) * (a + c) * (b + d))
    phi = 0.0 if denom == 0 else (a * d - b * c) / denom
    diff = p1 - p2
    dl = math.sqrt(max(0.0, (p1 - l1) ** 2 - 2 * phi * (p1 - l1) * (u2 - p2) + (u2 - p2) ** 2))
    du = math.sqrt(max(0.0, (u1 - p1) ** 2 - 2 * phi * (u1 - p1) * (p2 - l2) + (p2 - l2) ** 2))
    return round(100 * diff, 1), round(100 * (diff - dl), 1), round(100 * (diff + du), 1)


def kappa(x, y, levels, weighted=False):
    pairs = [(a, b) for a, b in zip(x, y) if a in levels and b in levels]
    n = len(pairs)
    if not n:
        return None
    k = len(levels)
    idx = {v: i for i, v in enumerate(levels)}
    w = (lambda i, j: 1 - abs(i - j) / (k - 1)) if weighted else (lambda i, j: 1.0 if i == j else 0.0)
    po = sum(w(idx[a], idx[b]) for a, b in pairs) / n
    px = [sum(1 for a, _ in pairs if a == v) / n for v in levels]
    py = [sum(1 for _, b in pairs if b == v) / n for v in levels]
    pe = sum(w(i, j) * px[i] * py[j] for i in range(k) for j in range(k))
    return round((po - pe) / (1 - pe), 3) if pe < 1 else 1.0


def read_tab(ws, key):
    """{(qid, system): (q1, q2)} from one rater tab. The number and id cells are merged over a question's three rows,
    so they are carried down; yellow example rows (id 'example-n') are skipped."""
    out, qid = {}, None
    for row in ws.iter_rows(min_row=2, values_only=True):
        row = tuple(row) + (None,) * (11 - len(row))  # a Google Sheets export drops empty trailing cells
        if row[1]:
            qid = str(row[1])
        letter = row[2]
        if not qid or not letter or qid.startswith('example'):
            continue
        system = key['letters'][qid][letter]
        out[(qid, system)] = tuple((row[i] or '').strip() if isinstance(row[i], str) else row[i] for i in (7, 8))
    return out


def binary(labels):
    q1, q2 = labels
    return {'correct': None if not q1 else q1 != 'Major error', 'correct_strict': None if not q1 else q1 == 'No error',
            'supported': None if q2 in (None, '', 'No data claims') else q2 == 'All supported'}


def answer_types(questions):
    """Automatic check of each system's answer type against the expected behaviour (first runs)."""
    out = {}
    exp = {q['id']: q['expected'] for q in questions}
    for system, run in RUNS.items():
        rows = {r['id']: r for r in map(json.loads, open(ROOT / '.cache/human-eval-20261020/runs' / run / 'results.jsonl'))}
        boundary = [q for q, e in exp.items() if e != 'answer']
        answerable = [q for q, e in exp.items() if e == 'answer']
        k1 = sum(rows[q]['answer_type'] in ('declined', 'limited') for q in boundary)
        k2 = sum(rows[q]['answer_type'] == 'answered' for q in answerable)
        out[system] = {'declined_or_limited_where_expected': {'k': k1, 'n': len(boundary), 'ci95': wilson(k1, len(boundary))},
                       'fully_answered_where_answerable': {'k': k2, 'n': len(answerable), 'ci95': wilson(k2, len(answerable))}}
    return out


def summarize(ratings, questions, subset=None):
    qids = [q['id'] for q in questions if subset is None or subset(q)]
    res = {}
    for s in SYSTEMS:
        res[s] = {}
        for crit in ('correct', 'correct_strict', 'supported'):
            vals = [binary(ratings[(q, s)])[crit] for q in qids if (q, s) in ratings]
            vals = [v for v in vals if v is not None]
            k, n = sum(vals), len(vals)
            res[s][crit] = {'k': k, 'n': n, 'pct': round(100 * k / n, 1) if n else None, 'ci95': wilson(k, n)}
        labels = [ratings[(q, s)][1] for q in qids if (q, s) in ratings]
        res[s]['q2_counts'] = {v: sum(1 for x in labels if x == v) for v in Q2}  # the four evidence labels
    for other in ('b1', 'b2'):
        for crit in ('correct', 'supported'):
            a = b = c = d = 0
            for q in qids:
                x, y = binary(ratings.get((q, 'askglp1'), ('', '')))[crit], binary(ratings.get((q, other), ('', '')))[crit]
                if x is None or y is None:
                    continue
                a += x and y; b += x and not y; c += (not x) and y; d += (not x) and (not y)
            res[f'askglp1_vs_{other}_{crit}'] = {'both': a, 'askglp1_only': b, f'{other}_only': c, 'neither': d,
                                                 'diff_pp_ci95': newcombe_paired(a, b, c, d), 'mcnemar_p': round(mcnemar_exact(b, c), 4)}
    return res


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--xlsx', required=True)
    ap.add_argument('--out')
    a = ap.parse_args(argv)
    key = json.loads(KEY.read_text())
    questions = json.loads(QUESTIONS.read_text())['questions']
    wb = load_workbook(a.xlsx, read_only=True)
    tabs = [n for n in wb.sheetnames if n != 'Instructions']
    ratings = {n: read_tab(wb[n], key) for n in tabs}
    report = {'answer_types': answer_types(questions), 'tabs': tabs, 'filled': {n: sum(1 for v in r.values() if all(v)) for n, r in ratings.items()}}
    # System results come from the Consensus tab (one row per answer, all 120 questions) once it exists; per-rater
    # summaries cover only that rater's questions.
    groups = {'all': None, 'to_answer': lambda q: q['expected'] == 'answer', 'to_decline_or_limit': lambda q: q['expected'] != 'answer',
              'faers': lambda q: q['area'] == 'rq2', 'not_faers': lambda q: q['area'] != 'rq2',
              'answerable_not_faers': lambda q: q['expected'] == 'answer' and q['area'] != 'rq2'}
    for n, r in ratings.items():
        report[n] = {g: summarize(r, questions, f) for g, f in groups.items()}
    # Agreement within each pair on that pair's questions.
    report['agreement'] = {}
    for pair, (r1, r2) in (key.get('pairs') or {}).items():
        if r1 not in ratings or r2 not in ratings:
            continue
        x, y = ratings[r1], ratings[r2]
        own = set((key.get('pair_questions') or {}).get(pair) or [])
        keys = sorted(k for k in set(x) & set(y) if k[0] in own)
        report['agreement'][f'{r1}+{r2}'] = {
            'answers': len(keys),
            'q1_weighted_kappa': kappa([x[k][0] for k in keys], [y[k][0] for k in keys], Q1, weighted=True),
            'q1_correct_kappa': kappa([str(binary(x[k])['correct']) for k in keys], [str(binary(y[k])['correct']) for k in keys], ['True', 'False']),
            'q2_kappa': kappa([x[k][1] for k in keys], [y[k][1] for k in keys], Q2)}
    text = json.dumps(report, indent=1)
    if a.out:
        Path(a.out).mkdir(parents=True, exist_ok=True)
        (Path(a.out) / 'scores.json').write_text(text)
    print(text)


if __name__ == '__main__':
    main()
