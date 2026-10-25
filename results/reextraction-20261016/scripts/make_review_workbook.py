"""Blind human-annotation workbook for the extraction test set (test300), in the style of the 2026-09-29 screening workbook.

One workbook: an Instructions tab (from human-validation/reviewer-guide.md) and one tab per reviewer, one row per post.
Reddit posts (200) are numbered 1-200 and go to the two Reddit reviewers; WebMD reviews (100) are numbered 201-300 and go
to the two WebMD reviewers. Posts are shuffled with a fixed seed. Reviewers see only the text: no system output, sampling
stratum or weight. Column titles are the questions; yellow example rows at the top of each tab show filled-in answers and
are not part of the sample (their Record ID starts with "example").

Writes .cache/reextraction-20261016/human-validation/{extraction-human-annotation.xlsx, key.jsonl, manifest.json}.
"""
import csv, hashlib, json, os, random, re

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.indexed_list import IndexedList
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
STAGE = os.path.join(ROOT, 'results/reextraction-20261016')
CACHE = os.path.join(ROOT, '.cache/reextraction-20261016')
SAMPLE = os.path.join(CACHE, 'samples/test300.jsonl')
WEBMD = os.path.join(ROOT, 'data/raw/webmd/2026-09-12/webmd_all_reviews.csv')
GUIDE = os.path.join(STAGE, 'human-validation/reviewer-guide.md')
OUT = os.path.join(CACHE, 'human-validation')
SEED = 20261018
REVIEWERS = {'reddit': ['Siddhartha', 'Anay'], 'webmd': ['Yeaeun', 'Haozheng']}
AGE_GROUPS = ['13-18', '19-24', '25-34', '35-44', '45-54', '55-64', '65-74', '75 or over', 'Not stated']
SEXES = ['female', 'male', 'other', 'Not stated']

# (column title, width, note shown on hover); the title says what to write
AGE = ('Age group of the writer (dropdown)', 13,
       'Only an age the writer states about themselves ("I\'m 52", "52F", "(45M)"). "Late 50s" is 55-64; "in my 50s" is Not stated.')
SEX = ('Sex of the writer (dropdown)', 12,
       'Only when the writer states it ("52F", "as a woman", "I\'m a guy"). Never guess from "my husband", pregnancy or menopause.')
DRUG = ('GLP-1 drug the writer takes (name, not named, or none)', 20,
        'The brand if the post names it (Ozempic, Wegovy, Rybelsus, Mounjaro, Zepbound, Victoza, Saxenda, Trulicity); otherwise '
        'the generic name if given; otherwise "not named". Write "none" if the post is not about the writer taking one of these drugs.')
REASONS = ('Reasons the writer gives for taking it (one per line, or none)', 28,
           'Only reasons stated with words like "for", "to", "because", "so", "prescribed for". Results ("lost 20 lbs", '
           '"helped my A1C") and goals without a reason word are not reasons.')
EFFECTS = ('Side effects the writer had from it (one per line, or none)', 28,
           'Symptoms the writer had and links to the drug, in the writer\'s words. Leave out symptoms that are feared, asked '
           'about, doubted, someone else\'s, or did not happen ("no nausea"). One symptom per line.')
DETAILS = ('Side-effect details (side effect: details, one per line)', 40,
           'Only details the writer gives: dose at the time, how long it lasted, mild/moderate/severe, kept from work or daily '
           'life, called or saw a doctor, ER, hospital, lowered/paused/skipped/stopped the drug. Example: "nausea: 5 mg, '
           'lasted 2 days, severe, went to the ER, stopped the drug".')
COMMENT = ('Comment (optional)', 30, None)

EXAMPLES = {
    'reddit': [
        ['Posted in r/zepbound\n\n32F here. Started Zepbound 2.5 three weeks ago to lose weight. The first week I had mild '
         'nausea, and I\'ve been constipated ever since. My sister is on it too and gets bad headaches.',
         '25-34', 'female', 'Zepbound', 'lose weight', 'nausea\nconstipated',
         'nausea: 2.5, lasted the first week, mild\nconstipated: 2.5', 'The headaches are the sister\'s, so they are left out.'],
        ['Posted in r/diabetes_t2\n\nMy doctor put me on Ozempic for my A1C. Not sure if it\'s the shot or a stomach bug, but '
         'I\'ve had diarrhea all week. Down 10 lbs so far!',
         'Not stated', 'Not stated', 'Ozempic', 'A1C', 'none', '',
         'The diarrhea is left out because the writer is not sure the drug caused it. Losing weight is not a side effect.'],
        ['Posted in r/mounjaro\n\nThe 7.5 dose made me so sick I couldn\'t keep anything down and I ended up in the ER. Went '
         'back down to 5 and it\'s better now.',
         'Not stated', 'Not stated', 'not named', 'none', 'sick',
         'sick: 7.5, couldn\'t keep anything down, went to the ER, lowered the dose', 'The post never names the drug.'],
    ],
    'webmd': [
        ['WebMD review of Trulicity (dulaglutide)\n\nMy doctor prescribed this for type 2 diabetes. I had severe stomach '
         'pain after every shot, so I stopped after a month.',
         'Trulicity', 'type 2 diabetes', 'stomach pain', 'stomach pain: severe, stopped the drug',
         '"After a month" is the time on the drug, not how long the pain lasted.'],
        ['WebMD review of Wegovy (semaglutide)\n\nWorks great! Lost 30 pounds and no side effects at all.',
         'Wegovy', 'none', 'none', '', 'Losing weight is a result, not a reason.'],
    ],
}

FONT, BOLD, TITLE = Font(name='Arial', size=10), Font(name='Arial', size=10, bold=True), Font(name='Arial', size=13, bold=True)
GRAY_FONT = Font(name='Arial', size=10, color='FF666666')
HEAD_FILL, ID_FILL = PatternFill('solid', fgColor='FFD9D9D9'), PatternFill('solid', fgColor='FFF3F3F3')
EXAMPLE_FILL = PatternFill('solid', fgColor='FFFFF2CC')
TOP_WRAP = Alignment(vertical='top', wrap_text=True)


def sha(path):
    return hashlib.sha256(open(path, 'rb').read()).hexdigest()


def record_text(rec, webmd_text):
    if rec['platform'] == 'reddit':
        return f"Posted in r/{rec['community']}\n\n{rec['text']}"
    p = rec['profile']
    return f"WebMD review of {p['brand']} ({p['ingredient']})\n\n{webmd_text[rec['record_id']]}"


def guide_rows(path):
    """Markdown guide -> instruction rows: (cells, style) with style in title/heading/table_head/table/text."""
    rows = []
    for raw in open(path, encoding='utf-8').read().splitlines():
        line = raw.strip()
        if not line:
            rows.append(([], 'text')); continue
        if re.fullmatch(r'\|[-|\s]+\|', line):
            continue
        clean = lambda s: s.replace('**', '').strip()
        if line.startswith('# '):
            rows.append(([clean(line[2:])], 'title'))
        elif line.startswith('## '):
            rows.append(([clean(line[3:])], 'heading'))
        elif line.startswith('|'):
            cells = [clean(c) for c in line.strip('|').split('|')]
            head = rows and rows[-1][1] not in ('table_head', 'table')
            rows.append((cells, 'table_head' if head else 'table'))
        elif line.startswith('- '):
            rows.append((['• ' + clean(line[2:])], 'text'))
        else:
            rows.append(([clean(line)], 'text'))
    out = []
    for cells, style in rows:  # collapse repeated blank lines; no blank line under a heading
        if not cells and out and (not out[-1][0] or out[-1][1] == 'heading'):
            continue
        out.append((cells, style))
    return out


def instructions(wb):
    ws = wb.active
    ws.title = 'Instructions'
    for col, width in zip('ABC', (95, 16, 70)):
        ws.column_dimensions[col].width = width
    for r, (cells, style) in enumerate(guide_rows(GUIDE), 1):
        for c, value in enumerate(cells, 1):
            cell = ws.cell(r, c, value)
            cell.alignment = TOP_WRAP
            cell.font = TITLE if style == 'title' else BOLD if style in ('heading', 'table_head') else FONT
            if style == 'table_head':
                cell.fill = HEAD_FILL


def dropdown(ws, col, last, values, error):
    dv = DataValidation(type='list', formula1='"' + ','.join(values) + '"', allow_blank=True, showDropDown=False,
                        showErrorMessage=True, error=error)
    dv.add(f'{col}2:{col}{last}')
    ws.add_data_validation(dv)


def reviewer_tab(wb, name, platform, records):
    ws = wb.create_sheet(name)
    reddit = platform == 'reddit'
    answer_cols = ([AGE, SEX] if reddit else []) + [DRUG, REASONS, EFFECTS, DETAILS, COMMENT]
    cols = [('No.', 7, None), ('Record ID', 20, None), ('Text', 80, None)] + answer_cols
    for c, (title, width, note) in enumerate(cols, 1):
        cell = ws.cell(1, c, title)
        cell.font, cell.fill, cell.alignment = BOLD, HEAD_FILL, TOP_WRAP
        if note:
            cell.comment = Comment(note, 'Instructions', width=320, height=140)
        ws.column_dimensions[cell.column_letter].width = width
    rows = [(f'Example {i}', f'example-{i}', ex, True) for i, ex in enumerate(EXAMPLES[platform], 1)]
    rows += [(rec['no'], rec['key'], [rec['text']], False) for rec in records]
    for r, (no, key, values, example) in enumerate(rows, 2):
        values = values + [None] * (len(cols) - 2 - len(values))
        for c in range(1, len(cols) + 1):
            cell = ws.cell(r, c, no if c == 1 else key if c == 2 else values[c - 3])
            cell.font, cell.alignment, cell.number_format = FONT, TOP_WRAP, '@'
            if c == 1:
                cell.font, cell.alignment, cell.number_format = GRAY_FONT, Alignment(horizontal='center', vertical='top', wrap_text=True), 'General'
            elif c == 2:
                cell.font, cell.alignment = GRAY_FONT, Alignment(vertical='top')
            if example:
                cell.fill = EXAMPLE_FILL
            elif c <= 2:
                cell.fill = ID_FILL
    last = len(rows) + 1
    letter = {title: ws.cell(1, c).column_letter for c, (title, _, _) in enumerate(cols, 1)}
    if reddit:
        dropdown(ws, letter[AGE[0]], last, AGE_GROUPS, 'Choose an age group or Not stated.')
        dropdown(ws, letter[SEX[0]], last, SEXES, 'Choose female, male, other or Not stated.')
    ws.freeze_panes = 'A2'


def main():
    sample = [json.loads(l) for l in open(SAMPLE)]
    want = {r['record_id'] for r in sample if r['platform'] == 'webmd'}
    webmd_text = {w['Review ID']: w['Textual Review'].strip() for w in csv.DictReader(open(WEBMD, encoding='utf-8')) if w['Review ID'] in want}
    rng = random.Random(SEED)
    records, key, no = {'reddit': [], 'webmd': []}, [], 0
    for platform in ('reddit', 'webmd'):
        recs = [r for r in sample if r['platform'] == platform]
        rng.shuffle(recs)
        for rec in recs:
            no += 1
            rk = f"{platform}:{rec['record_id']}"
            records[platform].append({'no': no, 'key': rk, 'text': record_text(rec, webmd_text)})
            key.append({'no': no, 'record': rk, 'platform': platform, 'stratum': rec['stratum'], 'weight': rec['weight'],
                        'reviewers': REVIEWERS[platform]})
    os.makedirs(OUT, exist_ok=True)
    wb = Workbook()
    wb._fonts = IndexedList([FONT])  # workbook default font Arial 10, as in the screening workbook
    instructions(wb)
    for platform in ('reddit', 'webmd'):
        for name in REVIEWERS[platform]:
            reviewer_tab(wb, name, platform, records[platform])
    path = os.path.join(OUT, 'extraction-human-annotation.xlsx')
    wb.save(path)
    with open(os.path.join(OUT, 'key.jsonl'), 'w') as f:
        for k in key:
            f.write(json.dumps(k, ensure_ascii=False) + '\n')
    manifest = {'design': 'blind: reviewers see only the post text; one row per post; yellow example rows on top',
                'workbook': os.path.relpath(path, ROOT), 'workbook_sha256': sha(path), 'key_sha256': sha(os.path.join(OUT, 'key.jsonl')),
                'sample': os.path.relpath(SAMPLE, ROOT), 'sample_sha256': sha(SAMPLE), 'guide': os.path.relpath(GUIDE, ROOT),
                'guide_sha256': sha(GUIDE), 'seed': SEED, 'reviewers': REVIEWERS, 'numbering': {'reddit': '1-200', 'webmd': '201-300'},
                'example_rows': {p: len(v) for p, v in EXAMPLES.items()}}
    json.dump(manifest, open(os.path.join(OUT, 'manifest.json'), 'w'), indent=1)
    print(json.dumps(manifest, indent=1))


if __name__ == '__main__':
    main()
